"""Tests del ``DispTOTParser`` (Totalizadores).

Cubre la extracción de ``Tabla_Disp_TOT`` (hoja ``DISP_TOT``) y la
construcción de ``DispTOT``. Los campos exclusivos ``tipo``/
``incxpulso``/``proceso`` y sus ``cfg_*`` se verifican explícitamente.
"""
from __future__ import annotations

from unittest.mock import MagicMock

from areas.alimentacion.data.data_dispositivos import DispTOT
from areas.alimentacion.helpers.excel.excel_parser_disp_tot import DispTOTParser
from core.infrastructure.config.config_manager import ConfigManager

from tests._disp_parser_test_helpers import (
    build_full_row,
    load_workbook_safe,
    save_xlsx_with_disp_table,
)


def test_extrae_disp_tot_basico(tmp_path) -> None:
    """1 fila con ``tipo``/``incxpulso``/``proceso`` populados."""
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "TOT",
        rows=[build_full_row("TOT", UID="TOT_001", Numero=1,
                              plc_tag="V_TOT_001",
                              **{"Tipo": "KG",
                                 "IncXPulso": 0.25,
                                 "Proceso": "PR2",
                                 "Cfg.Tipo": "cfg_tipo := 'KG';"})],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispTOTParser().extraer(wb)
    assert len(result) == 1
    d = result[0]
    assert isinstance(d, DispTOT)
    assert d.uid == "TOT_001"
    assert d.tipo == "KG"
    assert d.incxpulso == 0.25
    assert d.proceso == "PR2"
    assert d.cfg_tipo == "cfg_tipo := 'KG';"
    assert d.e_byte == 0
    assert d.e_bit == 0


def test_hoja_inexistente_devuelve_lista_vacia(tmp_path) -> None:
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "TOT",
        rows=[build_full_row("TOT")],
        sheet_name="OTRA_HOJA",
    )
    wb = load_workbook_safe(xlsx_path)
    assert DispTOTParser().extraer(wb) == []


def test_tabla_inexistente_devuelve_lista_vacia(tmp_path) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    xlsx_path = tmp_path / "otro.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "DISP_TOT"
    ws.append(["X", "Y"])
    ws.append([1, 2])
    table = Table(displayName="Tabla_Otra", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    wb.save(str(xlsx_path))

    wb2 = load_workbook_safe(xlsx_path)
    assert DispTOTParser().extraer(wb2) == []


def test_fila_sin_uid_se_descarta(tmp_path) -> None:
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "TOT",
        rows=[
            [None, None, "X", "X", "X"],
            build_full_row("TOT", UID="TOT_001"),
        ],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispTOTParser().extraer(wb)
    assert len(result) == 1
    assert result[0].uid == "TOT_001"


def test_defaults_when_only_uid_and_numero(tmp_path) -> None:
    """Solo UID+Numero -> ``tipo``/``incxpulso``/``proceso`` son defaults."""
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "TOT",
        rows=[["TOT_001", 1, "V_TOT_001", "c", "d"]],
        headers=["UID", "Numero", "PLC.Tag", "PLC.Comentario", "Descripcion"],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispTOTParser().extraer(wb)
    assert len(result) == 1
    d = result[0]
    assert d.tipo == ""
    assert d.incxpulso == 0.0
    assert d.proceso == ""
    assert d.cfg_tipo == ""


def test_parser_respeta_config_manager_override(tmp_path) -> None:
    """Si el config_manager sobreescribe sheet/table, el parser lo respeta."""
    cm = MagicMock(spec=ConfigManager)
    cm.get_excel_target_for.return_value = {
        "sheet": "OTRA_HOJA_TOT",
        "table": "Tabla_Disp_TOT_Alt",
        "canonical": "DispTOT",
    }

    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "TOT",
        rows=[build_full_row("TOT", UID="TOT_002")],
        sheet_name="OTRA_HOJA_TOT",
        table_name="Tabla_Disp_TOT_Alt",
    )
    wb = load_workbook_safe(xlsx_path)

    parser = DispTOTParser(config_manager=cm)
    assert parser.SHEET == "OTRA_HOJA_TOT"
    assert parser.TABLE == "Tabla_Disp_TOT_Alt"
    result = parser.extraer(wb)
    assert len(result) == 1
    assert result[0].uid == "TOT_002"
