"""Tests del ``DispMSINAParser`` (Motores Sinamics).

Cubre la extracción de ``Tabla_Disp_M_SINA`` (hoja ``DISP_M_SINA``)
y la construcción de ``DispMSINA``. Los campos exclusivos
``vel_min``/``vel_max``/``cons_k`` y sus ``cfg_*`` se verifican
explícitamente.
"""
from __future__ import annotations

from areas.alimentacion.data.data_dispositivos import DispMSINA
from areas.alimentacion.helpers.excel.excel_parser_disp_m_sina import (
    DispMSINAParser,
)

from tests._disp_parser_test_helpers import (
    build_full_row,
    load_workbook_safe,
    save_xlsx_with_disp_table,
)


def test_extrae_disp_m_sina_basico(tmp_path) -> None:
    """1 fila con ``vel_min``/``vel_max``/``cons_k`` populados."""
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "M_SINA",
        rows=[build_full_row("M_SINA", UID="MSINA_001", Numero=1,
                              plc_tag="V_MSINA_001",
                              **{"Vel.Min": 10.5, "Vel.Max": 50.0,
                                 "ConsK": 0.85,
                                 "Cfg.VelMin": "cfg_velmin := 10.5;",
                                 "Cfg.VelMax": "cfg_velmax := 50.0;",
                                 "Cfg.ConstK": "cfg_constk := 0.85;"})],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispMSINAParser().extraer(wb)
    assert len(result) == 1
    d = result[0]
    assert isinstance(d, DispMSINA)
    assert d.uid == "MSINA_001"
    assert d.vel_min == 10.5
    assert d.vel_max == 50.0
    assert d.cons_k == 0.85
    assert d.cfg_vel_min == "cfg_velmin := 10.5;"
    assert d.cfg_vel_max == "cfg_velmax := 50.0;"
    assert d.cfg_cons_k == "cfg_constk := 0.85;"


def test_hoja_inexistente_devuelve_lista_vacia(tmp_path) -> None:
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "M_SINA",
        rows=[build_full_row("M_SINA")],
        sheet_name="OTRA_HOJA",
    )
    wb = load_workbook_safe(xlsx_path)
    assert DispMSINAParser().extraer(wb) == []


def test_tabla_inexistente_devuelve_lista_vacia(tmp_path) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    xlsx_path = tmp_path / "otro.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "DISP_M_SINA"
    ws.append(["X", "Y"])
    ws.append([1, 2])
    table = Table(displayName="Tabla_Otra", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    wb.save(str(xlsx_path))

    wb2 = load_workbook_safe(xlsx_path)
    assert DispMSINAParser().extraer(wb2) == []


def test_fila_sin_uid_se_descarta(tmp_path) -> None:
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "M_SINA",
        rows=[
            [None, None, "X", "X", "X"],
            build_full_row("M_SINA", UID="MSINA_001"),
        ],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispMSINAParser().extraer(wb)
    assert len(result) == 1
    assert result[0].uid == "MSINA_001"


def test_defaults_when_only_uid_and_numero(tmp_path) -> None:
    """Solo UID+Numero -> ``vel_min``/``vel_max``/``cons_k`` son ``0.0``."""
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "M_SINA",
        rows=[["MSINA_001", 1, "V_MSINA_001", "c", "d"]],
        headers=["UID", "Numero", "PLC.Tag", "PLC.Comentario", "Descripcion"],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispMSINAParser().extraer(wb)
    assert len(result) == 1
    d = result[0]
    assert d.vel_min == 0.0
    assert d.vel_max == 0.0
    assert d.cons_k == 0.0
    assert d.cfg_vel_min == ""
    assert d.cfg_vel_max == ""
    assert d.cfg_cons_k == ""


def test_parser_respeta_config_manager_override(tmp_path) -> None:
    """Si el config_manager sobreescribe sheet/table, el parser lo respeta."""
    from core.infrastructure.config.config_manager import ConfigManager
    from unittest.mock import MagicMock

    cm = MagicMock(spec=ConfigManager)
    cm.get_excel_target_for.return_value = {
        "sheet": "OTRA_HOJA_M_SINA",
        "table": "Tabla_Disp_M_SINA_Alt",
        "canonical": "DispM_SINA",
    }

    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "M_SINA",
        rows=[build_full_row("M_SINA", UID="MSINA_002")],
        sheet_name="OTRA_HOJA_M_SINA",
        table_name="Tabla_Disp_M_SINA_Alt",
    )
    wb = load_workbook_safe(xlsx_path)

    parser = DispMSINAParser(config_manager=cm)
    assert parser.SHEET == "OTRA_HOJA_M_SINA"
    assert parser.TABLE == "Tabla_Disp_M_SINA_Alt"
    result = parser.extraer(wb)
    assert len(result) == 1
    assert result[0].uid == "MSINA_002"
