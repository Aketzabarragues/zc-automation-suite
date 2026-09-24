"""Tests del ``DispPIDParser`` (Controladores PID).

Cubre la extracción de ``Tabla_Disp_PID`` (hoja ``DISP_PID``) y la
construcción de ``DispPID``. Los campos exclusivos ``proceso``/``pv``/
``disp_tipo``/``disp_tag`` se verifican explícitamente.
"""
from __future__ import annotations

from unittest.mock import MagicMock

from areas.alimentacion.data.data_dispositivos import DispPID
from areas.alimentacion.helpers.excel.excel_parser_disp_pid import DispPIDParser
from core.infrastructure.config.config_manager import ConfigManager

from tests._disp_parser_test_helpers import (
    build_full_row,
    load_workbook_safe,
    save_xlsx_with_disp_table,
)


def test_extrae_disp_pid_basico(tmp_path) -> None:
    """1 fila con ``proceso``/``pv``/``disp_tipo``/``disp_tag`` populados."""
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "PID",
        rows=[build_full_row("PID", UID="PID_001", Numero=1,
                              plc_tag="V_PID_001",
                              **{"Proceso": "PR3",
                                 "PV": "V_TEMP_001",
                                 "Disp.Tipo": "TEND",
                                 "Disp.Tag": "TEMP_PV"})],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispPIDParser().extraer(wb)
    assert len(result) == 1
    d = result[0]
    assert isinstance(d, DispPID)
    assert d.uid == "PID_001"
    assert d.proceso == "PR3"
    assert d.pv == "V_TEMP_001"
    assert d.disp_tipo == "TEND"
    assert d.disp_tag == "TEMP_PV"
    # NO tiene campos digitales ni analogicos.
    assert d.plc_index == 0
    assert d.hmi_index == 0


def test_hoja_inexistente_devuelve_lista_vacia(tmp_path) -> None:
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "PID",
        rows=[build_full_row("PID")],
        sheet_name="OTRA_HOJA",
    )
    wb = load_workbook_safe(xlsx_path)
    assert DispPIDParser().extraer(wb) == []


def test_tabla_inexistente_devuelve_lista_vacia(tmp_path) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    xlsx_path = tmp_path / "otro.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "DISP_PID"
    ws.append(["X", "Y"])
    ws.append([1, 2])
    table = Table(displayName="Tabla_Otra", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    wb.save(str(xlsx_path))

    wb2 = load_workbook_safe(xlsx_path)
    assert DispPIDParser().extraer(wb2) == []


def test_fila_sin_uid_se_descarta(tmp_path) -> None:
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "PID",
        rows=[
            [None, None, "X", "X", "X"],
            build_full_row("PID", UID="PID_001"),
        ],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispPIDParser().extraer(wb)
    assert len(result) == 1
    assert result[0].uid == "PID_001"


def test_defaults_when_only_uid_and_numero(tmp_path) -> None:
    """Solo UID+Numero -> ``proceso``/``pv``/``disp_tipo``/``disp_tag`` son ``""``."""
    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "PID",
        rows=[["PID_001", 1, "V_PID_001", "c", "d"]],
        headers=["UID", "Numero", "PLC.Tag", "PLC.Comentario", "Descripcion"],
    )
    wb = load_workbook_safe(xlsx_path)
    result = DispPIDParser().extraer(wb)
    assert len(result) == 1
    d = result[0]
    assert d.proceso == ""
    assert d.pv == ""
    assert d.disp_tipo == ""
    assert d.disp_tag == ""


def test_parser_respeta_config_manager_override(tmp_path) -> None:
    """Si el config_manager sobreescribe sheet/table, el parser lo respeta."""
    cm = MagicMock(spec=ConfigManager)
    cm.get_excel_target_for.return_value = {
        "sheet": "OTRA_HOJA_PID",
        "table": "Tabla_Disp_PID_Alt",
        "canonical": "DispPID",
    }

    xlsx_path = save_xlsx_with_disp_table(
        tmp_path, "PID",
        rows=[build_full_row("PID", UID="PID_002")],
        sheet_name="OTRA_HOJA_PID",
        table_name="Tabla_Disp_PID_Alt",
    )
    wb = load_workbook_safe(xlsx_path)

    parser = DispPIDParser(config_manager=cm)
    assert parser.SHEET == "OTRA_HOJA_PID"
    assert parser.TABLE == "Tabla_Disp_PID_Alt"
    result = parser.extraer(wb)
    assert len(result) == 1
    assert result[0].uid == "PID_002"
